import os
import sys
import json
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt, Inches, RGBColor

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

OUTPUT_DIR = "data/outputs"
NOTE_SPEAKER = ("Speaker = the person talking (\"I will...\"). The transcript has no "
                "speaker labels, so the speaker's name is not known.")
FOOTER = ("Generated automatically by the AI Meeting Minutes Generator. "
          "Please review before sharing.")

# ---------- helpers ----------
def _owners(a):
    return ", ".join(a["owners"]) if a["owners"] else "Unassigned"

def _deadline(a):
    if not a.get("deadline"):
        return "-"
    if a.get("deadline_date"):
        return f"{a['deadline']} ({a['deadline_date']})"
    return a["deadline"]

def _next(n):
    text = n["sentence"]
    if n.get("when"):
        text += f" - {n['when']}"
        if n.get("date"):
            text += f" ({n['date']})"
    return text

# ---------- DOCX ----------
def export_docx(m, out_path):
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    doc.add_heading(m["title"], level=0)
    meta = doc.add_paragraph()
    meta.add_run("Date: ").bold = True
    meta.add_run(m["date"])
    meta.add_run("      Source: ").bold = True
    meta.add_run(m["source_file"])

    doc.add_heading("Participants", level=1)
    for p in m["participants"] or ["(not identified)"]:
        doc.add_paragraph(p, style="List Bullet")

    doc.add_heading("Agenda / Topics Discussed", level=1)
    for t in m["agenda"] or ["(none identified)"]:
        doc.add_paragraph(t, style="List Number")

    doc.add_heading("Discussion Summary", level=1)
    doc.add_paragraph(m["summary"] or "(no summary generated)")

    doc.add_heading("Key Decisions", level=1)
    for d in m["key_decisions"] or ["(none found)"]:
        doc.add_paragraph(d, style="List Bullet")

    doc.add_heading("Action Items", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, h in enumerate(["#", "Task", "Owner", "Deadline"]):
        run = table.rows[0].cells[i].paragraphs[0].add_run(h)
        run.bold = True
    for i, a in enumerate(m["action_items"], 1):
        row = table.add_row().cells
        row[0].text = str(i)
        row[1].text = a["task"]
        row[2].text = _owners(a)
        row[3].text = _deadline(a)
    table.autofit = False
    widths = [Inches(0.4), Inches(3.8), Inches(1.2), Inches(1.4)]
    for row in table.rows:
        for idx, w in enumerate(widths):
            row.cells[idx].width = w
    note = doc.add_paragraph()
    r = note.add_run(NOTE_SPEAKER)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_heading("Next Meeting", level=1)
    for n in m["next_meeting"] or [{"sentence": "(not mentioned)"}]:
        doc.add_paragraph(_next(n), style="List Bullet")

    foot = doc.add_paragraph()
    r = foot.add_run(FOOTER)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.save(out_path)

# ---------- PDF ----------
def export_pdf(m, out_path):
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1x", parent=styles["Heading2"],
                        textColor=colors.HexColor("#1F3A5F"), spaceBefore=12, spaceAfter=4)
    body = ParagraphStyle("Bodyx", parent=styles["BodyText"], fontSize=10.5, leading=14)
    item = ParagraphStyle("Itemx", parent=body, leftIndent=16, bulletIndent=3)
    small = ParagraphStyle("Smallx", parent=body, fontSize=8.5, leading=11,
                           textColor=colors.HexColor("#666666"))
    cell = ParagraphStyle("Cellx", parent=body, fontSize=9.5, leading=12)

    doc = SimpleDocTemplate(out_path, pagesize=A4, title=m["title"],
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    story = [
        Paragraph(escape(m["title"]), styles["Title"]),
        Paragraph(f"<b>Date:</b> {escape(m['date'])} &nbsp;&nbsp;&nbsp; "
                  f"<b>Source:</b> {escape(m['source_file'])}", body),
    ]

    def bullets(items, empty):
        for t in items or [empty]:
            story.append(Paragraph(escape(t), item, bulletText="\u2022"))

    story.append(Paragraph("Participants", h1))
    bullets(m["participants"], "(not identified)")

    story.append(Paragraph("Agenda / Topics Discussed", h1))
    for i, t in enumerate(m["agenda"] or ["(none identified)"], 1):
        story.append(Paragraph(escape(t), item, bulletText=f"{i}."))

    story.append(Paragraph("Discussion Summary", h1))
    story.append(Paragraph(escape(m["summary"] or "(no summary generated)"), body))

    story.append(Paragraph("Key Decisions", h1))
    bullets(m["key_decisions"], "(none found)")

    story.append(Paragraph("Action Items", h1))
    data = [[Paragraph("<b>#</b>", cell), Paragraph("<b>Task</b>", cell),
             Paragraph("<b>Owner</b>", cell), Paragraph("<b>Deadline</b>", cell)]]
    for i, a in enumerate(m["action_items"], 1):
        data.append([Paragraph(str(i), cell), Paragraph(escape(a["task"]), cell),
                     Paragraph(escape(_owners(a)), cell),
                     Paragraph(escape(_deadline(a)), cell)])
    table = Table(data, colWidths=[1 * cm, 9 * cm, 3.5 * cm, 3.5 * cm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F2")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)
    story.append(Spacer(1, 4))
    story.append(Paragraph(escape(NOTE_SPEAKER), small))

    story.append(Paragraph("Next Meeting", h1))
    bullets([_next(n) for n in m["next_meeting"]], "(not mentioned)")

    story.append(Spacer(1, 14))
    story.append(Paragraph(escape(FOOTER), small))
    doc.build(story)

# ---------- both formats ----------
def export_all(minutes, out_dir=OUTPUT_DIR):
    os.makedirs(out_dir, exist_ok=True)
    stem = os.path.splitext(minutes["source_file"])[0]
    docx_path = os.path.join(out_dir, f"{stem}_minutes.docx")
    pdf_path = os.path.join(out_dir, f"{stem}_minutes.pdf")
    export_docx(minutes, docx_path)
    export_pdf(minutes, pdf_path)
    return docx_path, pdf_path

if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8-sig") as f:   # a *_minutes.json file
        minutes = json.load(f)
    d, p = export_all(minutes)
    print("Saved:", d)
    print("Saved:", p)